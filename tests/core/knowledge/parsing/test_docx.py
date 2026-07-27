"""DOCX L1 parsing keeps document order and embedded-media evidence."""

from __future__ import annotations

from io import BytesIO

from docx import Document
from PIL import Image

from core.knowledge.parsing import DocxParser, ParseRequest


def _docx_payload() -> bytes:
    document = Document()
    document.core_properties.title = "Sage Retrieval Review"
    document.add_heading("Retrieval trade-offs", level=1)
    document.add_paragraph("Exact scan remains the baseline.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Route"
    table.cell(0, 1).text = "P95"
    table.cell(1, 0).text = "exact"
    table.cell(1, 1).text = "36 ms"
    image_bytes = BytesIO()
    Image.new("RGB", (24, 12), "white").save(image_bytes, format="PNG")
    document.add_picture(BytesIO(image_bytes.getvalue()))
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_docx_parser_preserves_paragraph_table_and_embedded_image_evidence() -> None:
    request = ParseRequest(
        source_id="src_docx",
        relative_path="reports/retrieval-review.docx",
        source_revision="sha256:docx",
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        payload=_docx_payload(),
    )

    document = DocxParser().parse(request)

    assert document.title == "Sage Retrieval Review"
    assert document.provenance.parser_id == "sage.docx"
    assert document.provenance.parser_version == "1.0.0"
    assert [block.kind for block in document.blocks] == [
        "heading",
        "paragraph",
        "table",
        "media",
    ]
    assert document.blocks[1].heading_path == ("Retrieval trade-offs",)
    assert "| Route | P95 |" in document.blocks[2].text
    assert "| exact | 36 ms |" in document.rendered_markdown
    media = document.blocks[3]
    assert media.media_ref == "word/media/image1.png"
    assert media.confidence == 1.0
    assert media.page is None
    assert media.bbox is None
    assert DocxParser().parse(request) == document
