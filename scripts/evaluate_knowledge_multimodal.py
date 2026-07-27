#!/usr/bin/env python3
"""Run the deterministic PR-7 L1/L2 multimodal evidence contract evaluation."""

from __future__ import annotations

import argparse
import asyncio
import hashlib
import json
import subprocess
import time
import zipfile
from io import BytesIO
from pathlib import Path
from statistics import median
from typing import Any

import httpx
from docx import Document
from PIL import Image, PngImagePlugin

from core.knowledge.multimodal_eval import (
    MultimodalEvalCase,
    load_multimodal_cases,
    multimodal_cases_sha256,
)
from core.knowledge.parsing import DocxParser, ParseRequest, PngImageParser
from core.knowledge.parsing.adapters.qwen_vl import QwenVlAdapter, QwenVlConfig
from core.knowledge.retrieval import chunk_document, citation_id

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CASES = ROOT / "knowledge/eval/multimodal_cases.jsonl"
DEFAULT_OUTPUT = ROOT / "evals/reports/knowledge_multimodal_evidence_v1_2026-07-28.json"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--cases", type=Path, default=DEFAULT_CASES)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--allow-dirty", action="store_true")
    args = parser.parse_args()

    dirty = bool(_git("status", "--porcelain"))
    if dirty and not args.allow_dirty:
        raise SystemExit("refusing formal multimodal report from dirty source")
    cases = load_multimodal_cases(args.cases)
    report = evaluate(cases, cases_path=args.cases, source_dirty=dirty)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    output_label = (
        str(args.output.resolve().relative_to(ROOT))
        if args.output.resolve().is_relative_to(ROOT)
        else str(args.output.resolve())
    )
    print(
        json.dumps(
            {
                "output": output_label,
                "case_count": len(cases),
                "overall_passed": report["gates"]["overall_passed"],
                "deterministic_digest": report["deterministic_digest"],
            },
            ensure_ascii=False,
        )
    )


def evaluate(
    cases: tuple[MultimodalEvalCase, ...],
    *,
    cases_path: Path,
    source_dirty: bool,
) -> dict[str, Any]:
    documents: dict[str, Any] = {}
    case_results: list[dict[str, Any]] = []
    latencies: list[float] = []
    citation_by_block: dict[tuple[str, str], str] = {}
    citation_stable = True
    for case in cases:
        started = time.perf_counter()
        document = documents.get(case.fixture_id)
        if document is None:
            document = _document(case.fixture_id)
            documents[case.fixture_id] = document
        block = next(
            (
                item
                for item in document.blocks
                if item.kind == case.expected_block_kind and case.required_text in item.text
            ),
            None,
        )
        checks = {
            "required_block": block is not None,
            "parser": (
                document.provenance.parser_id == case.expected_parser_id
                and document.provenance.parser_version == case.expected_parser_version
            ),
        }
        located_citation = None
        if block is not None:
            chunks = chunk_document(document, **_chunk_kwargs(case))
            chunk = next((item for item in chunks if item.block_id == block.block_id), None)
            checks.update(
                {
                    "chunk_projection": chunk is not None,
                    "page": block.page == case.gold_page,
                    "bbox": block.bbox == case.gold_bbox,
                    "media_ref": block.media_ref == case.expected_media_ref,
                    "confidence": 0.0 <= block.confidence <= 1.0,
                }
            )
            if chunk is not None:
                located_citation = citation_id(chunk)
                block_key = (case.fixture_id, block.block_id)
                previous = citation_by_block.setdefault(block_key, located_citation)
                citation_stable = citation_stable and previous == located_citation
                checks["citation_projection"] = (
                    chunk.page_number == block.page
                    and chunk.bbox == block.bbox
                    and chunk.media_ref == block.media_ref
                    and chunk.confidence == block.confidence
                    and chunk.parser_id == document.provenance.parser_id
                    and chunk.parser_version == document.provenance.parser_version
                )
        passed = all(checks.values())
        latency_ms = (time.perf_counter() - started) * 1_000
        latencies.append(latency_ms)
        case_results.append(
            {
                "case_id": case.case_id,
                "dataset_split": case.dataset_split,
                "fixture_id": case.fixture_id,
                "layer": case.layer,
                "passed": passed,
                "checks": checks,
                "citation_id": located_citation,
                "latency_ms": round(latency_ms, 3),
            }
        )

    visual = [item for item, case in zip(case_results, cases, strict=True) if case.gold_bbox]
    test = [item for item in case_results if item["dataset_split"] == "test"]
    pass_rate = _ratio(sum(bool(item["passed"]) for item in case_results), len(case_results))
    bbox_accuracy = _ratio(
        sum(bool(item["checks"].get("bbox")) for item in visual),
        len(visual),
    )
    test_pass_rate = _ratio(sum(bool(item["passed"]) for item in test), len(test))
    metrics = {
        "case_pass_rate": pass_rate,
        "test_case_pass_rate": test_pass_rate,
        "visual_bbox_accuracy": bbox_accuracy,
        "citation_identity_stable": citation_stable,
        "latency_ms": {
            "p50": round(median(latencies), 3),
            "p95": round(_percentile(latencies, 0.95), 3),
        },
    }
    gates: dict[str, Any] = {
        "case_pass_rate": {"minimum": 1.0, "actual": pass_rate, "passed": pass_rate == 1.0},
        "frozen_test_pass_rate": {
            "minimum": 1.0,
            "actual": test_pass_rate,
            "passed": test_pass_rate == 1.0,
        },
        "visual_bbox_accuracy": {
            "minimum": 1.0,
            "actual": bbox_accuracy,
            "passed": bbox_accuracy == 1.0,
        },
        "citation_identity_stable": {"passed": citation_stable},
    }
    gates["overall_passed"] = all(bool(value["passed"]) for value in gates.values())
    deterministic_cases = [
        {key: value for key, value in item.items() if key != "latency_ms"} for item in case_results
    ]
    deterministic_metrics = {key: value for key, value in metrics.items() if key != "latency_ms"}
    deterministic = {
        "cases_sha256": multimodal_cases_sha256(cases_path),
        "metrics": deterministic_metrics,
        "cases": deterministic_cases,
    }
    digest = (
        "sha256:"
        + hashlib.sha256(
            json.dumps(
                deterministic, ensure_ascii=False, sort_keys=True, separators=(",", ":")
            ).encode()
        ).hexdigest()
    )
    return {
        "schema_version": 1,
        "evaluation_id": "sage-knowledge-multimodal-evidence-v1",
        "inputs": {
            "cases": str(cases_path.resolve().relative_to(ROOT)),
            "cases_sha256": multimodal_cases_sha256(cases_path),
            "case_count": len(cases),
            "split_counts": {
                split: sum(case.dataset_split == split for case in cases)
                for split in ("dev", "calibration", "test")
            },
        },
        "source": {"commit": _git("rev-parse", "HEAD"), "dirty": source_dirty},
        "scope": {
            "l1_local_parsers": ["sage.docx@1.0.0", "sage.png@1.0.0"],
            "l2_vlm_contract": "qwen3-vl@2.0.0 fixture response",
            "live_vlm_quality_evaluated": False,
            "visual_vector_retrieval_enabled": False,
            "fixture_origin": "project-authored synthetic documents and responses",
        },
        "metrics": metrics,
        "gates": gates,
        "cases": case_results,
        "deterministic_digest": digest,
    }


def _document(fixture_id: str) -> Any:
    if fixture_id == "docx_ordered":
        payload = _docx_payload()
        return DocxParser().parse(
            _request(
                "report.docx",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                payload,
            )
        )
    if fixture_id in {"png_metadata", "png_no_description"}:
        return PngImageParser().parse(
            _request(
                "retrieval.png",
                "image/png",
                _png_payload(with_metadata=fixture_id == "png_metadata"),
            )
        )
    if fixture_id == "qwen_region":
        return asyncio.run(_qwen_document())
    raise ValueError(f"unknown multimodal fixture: {fixture_id}")


def _docx_payload() -> bytes:
    document = Document()
    document.core_properties.title = "Sage Retrieval Review"
    document.add_heading("Retrieval trade-offs", level=1)
    document.add_paragraph("Exact scan remains the baseline.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Route"
    table.cell(0, 1).text = "P95"
    table.cell(1, 0).text = "exact"
    table.cell(1, 1).text = "36 ms"
    image = BytesIO()
    Image.new("RGB", (24, 12), "white").save(image, format="PNG")
    document.add_picture(BytesIO(image.getvalue()))
    output = BytesIO()
    document.save(output)
    return _deterministic_docx(output.getvalue())


def _deterministic_docx(payload: bytes) -> bytes:
    source = BytesIO(payload)
    output = BytesIO()
    with (
        zipfile.ZipFile(source) as archive,
        zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as normalized,
    ):
        for name in sorted(archive.namelist()):
            original = archive.getinfo(name)
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = original.external_attr
            normalized.writestr(info, archive.read(name))
    return output.getvalue()


def _png_payload(*, with_metadata: bool) -> bytes:
    output = BytesIO()
    metadata = None
    if with_metadata:
        metadata = PngImagePlugin.PngInfo()
        metadata.add_text("Title", "Exact retrieval chart")
        metadata.add_text("Description", "PostgreSQL exact scan P95 is 36 milliseconds")
    Image.new("RGB", (320, 180), "white").save(output, format="PNG", pnginfo=metadata)
    return output.getvalue()


async def _qwen_document() -> Any:
    def handler(_request_value: httpx.Request) -> httpx.Response:
        return httpx.Response(
            200,
            json={
                "choices": [
                    {
                        "message": {
                            "content": json.dumps(
                                {
                                    "title": "Retrieval chart",
                                    "regions": [
                                        {
                                            "kind": "table",
                                            "text": "exact P95 36 ms; cross encoder P95 1436 ms",
                                            "bbox": [0.1, 0.2, 0.9, 0.8],
                                            "confidence": 0.93,
                                        }
                                    ],
                                }
                            )
                        }
                    }
                ]
            },
        )

    async def progress(_value: Any) -> None:
        return None

    async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
        adapter = QwenVlAdapter(QwenVlConfig(api_key="fixture-key"), client=client)
        return await adapter.parse(
            _request("charts/retrieval.png", "image/png", _png_payload(with_metadata=False)),
            progress=progress,
        )


def _request(relative_path: str, media_type: str, payload: bytes) -> ParseRequest:
    return ParseRequest(
        source_id="src_multimodal_fixture",
        relative_path=relative_path,
        source_revision="sha256:" + hashlib.sha256(payload).hexdigest(),
        media_type=media_type,
        payload=payload,
    )


def _chunk_kwargs(case: MultimodalEvalCase) -> dict[str, Any]:
    return {
        "workspace_id": "knowledge-multimodal-eval",
        "page_id": f"page_{case.fixture_id}",
        "page_revision": f"krev_{case.fixture_id}",
        "page_path": f"wiki/fixtures/{case.fixture_id}.md",
        "source_id": f"src_{case.fixture_id}",
        "source_revision": f"sha256:{case.fixture_id}",
        "source_kind": "eval_fixture",
        "source_relative_path": f"fixtures/{case.fixture_id}",
        "proposal_id": f"kprop_{case.fixture_id}",
        "artifact_id": f"part_{case.fixture_id}",
        "title": case.fixture_id,
        "visibility": "private",
        "active": True,
    }


def _percentile(values: list[float], ratio: float) -> float:
    ordered = sorted(values)
    index = max(0, min(len(ordered) - 1, int((len(ordered) - 1) * ratio + 0.999999)))
    return ordered[index]


def _ratio(numerator: int, denominator: int) -> float:
    return round(numerator / denominator, 6) if denominator else 0.0


def _git(*args: str) -> str:
    return subprocess.run(
        ["git", *args],
        cwd=ROOT,
        check=True,
        capture_output=True,
        text=True,
    ).stdout.strip()


if __name__ == "__main__":
    main()
